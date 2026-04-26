"""Orchestrates Excel workbook creation for a single deal."""
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from config import OUTPUTS_DIR
from models.property import Property
from models.underwriting import UWAssumptions, UWResults
from excel.sheets import (
    build_summary, build_assumptions, build_income,
    build_opex, build_debt_schedule, build_returns, build_sensitivity,
)


def _safe_filename(s: str) -> str:
    return re.sub(r"[^\w\s-]", "", s).strip().replace(" ", "_")[:40]


def generate_excel(
    prop: Property,
    assumptions: UWAssumptions,
    results: UWResults,
    output_dir: Path = OUTPUTS_DIR,
) -> Path:
    wb = Workbook()

    tabs = [
        ("1. Summary", build_summary),
        ("2. Assumptions", build_assumptions),
        ("3. Income", build_income),
        ("4. OpEx", build_opex),
        ("5. Debt Schedule", build_debt_schedule),
        ("6. Returns", build_returns),
        ("7. Sensitivity", build_sensitivity),
    ]

    # Use default sheet for first tab
    first_ws = wb.active
    first_ws.title = tabs[0][0]
    tabs[0][1](first_ws, prop, assumptions, results)

    for name, builder in tabs[1:]:
        ws = wb.create_sheet(title=name)
        if name in ("3. Income", "4. OpEx", "5. Debt Schedule", "6. Returns"):
            builder(ws, prop, assumptions, results)
        elif name == "2. Assumptions":
            builder(ws, prop, assumptions)
        elif name == "7. Sensitivity":
            builder(ws, prop, assumptions, results)

    # Tab colors
    tab_colors = ["1F3864", "2E75B6", "2E75B6", "2E75B6", "2E75B6", "375623", "843C0C"]
    for i, (ws, color) in enumerate(zip(wb.worksheets, tab_colors)):
        ws.sheet_properties.tabColor = color

    # Freeze panes on returns sheet
    returns_ws = wb["6. Returns"]
    returns_ws.freeze_panes = "B4"

    fname = f"{_safe_filename(prop.address)}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    output_path = output_dir / fname
    wb.save(str(output_path))
    return output_path
