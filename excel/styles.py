"""openpyxl style constants for PE-quality Excel output."""
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)

# Colors
NAVY = "1F3864"
LIGHT_BLUE = "DCE6F1"
MID_BLUE = "B8CCE4"
WHITE = "FFFFFF"
YELLOW = "FFFF99"
GREEN = "E2EFDA"
RED_LIGHT = "FCE4D6"
GRAY_LIGHT = "F2F2F2"
GRAY_MID = "D9D9D9"
BLACK = "000000"

# Fills
FILL_HEADER = PatternFill("solid", fgColor=NAVY)
FILL_SUBHEADER = PatternFill("solid", fgColor=MID_BLUE)
FILL_ALT = PatternFill("solid", fgColor=LIGHT_BLUE)
FILL_INPUT = PatternFill("solid", fgColor=YELLOW)
FILL_TOTAL = PatternFill("solid", fgColor=GRAY_MID)
FILL_GREEN = PatternFill("solid", fgColor=GREEN)
FILL_RED = PatternFill("solid", fgColor=RED_LIGHT)
FILL_GRAY = PatternFill("solid", fgColor=GRAY_LIGHT)

# Fonts
FONT_HEADER = Font(name="Calibri", bold=True, color=WHITE, size=11)
FONT_SUBHEADER = Font(name="Calibri", bold=True, color=NAVY, size=10)
FONT_BODY = Font(name="Calibri", size=10)
FONT_TOTAL = Font(name="Calibri", bold=True, size=10)
FONT_TITLE = Font(name="Calibri", bold=True, color=NAVY, size=14)
FONT_METRIC = Font(name="Calibri", bold=True, size=12)
FONT_SMALL = Font(name="Calibri", size=9)

# Borders
_THIN = Side(style="thin")
_MED = Side(style="medium")
BORDER_ALL = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
BORDER_BOTTOM = Border(bottom=_THIN)
BORDER_THICK_BOTTOM = Border(bottom=_MED)
BORDER_TOP_BOTTOM = Border(top=_THIN, bottom=_THIN)

# Alignments
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_WRAP = Alignment(wrap_text=True, vertical="center")

# Number formats
FMT_DOLLAR = '$#,##0'
FMT_DOLLAR_DEC = '$#,##0.00'
FMT_PCT = '0.00%'
FMT_PCT_1 = '0.0%'
FMT_NUM = '#,##0'
FMT_NUM_DEC = '#,##0.00'
FMT_MULT = '0.00"x"'


def style_header_row(ws, row: int, min_col: int, max_col: int):
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL


def style_subheader_row(ws, row: int, min_col: int, max_col: int):
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = FILL_SUBHEADER
        cell.font = FONT_SUBHEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL


def style_total_row(ws, row: int, min_col: int, max_col: int):
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = FILL_TOTAL
        cell.font = FONT_TOTAL
        cell.border = BORDER_ALL


def style_alt_row(ws, row: int, min_col: int, max_col: int, alt: bool):
    fill = FILL_ALT if alt else None
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row=row, column=col)
        if fill:
            cell.fill = fill
        cell.font = FONT_BODY
        cell.border = BORDER_ALL


def apply_cell(ws, row: int, col: int, value, fmt: str = None,
               fill=None, font=None, align=None, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    return cell
